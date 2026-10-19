import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A
    C_NAVY_CARD = RGBColor(30, 41, 59)      # #1E293B
    C_BLUE_PRIMARY = RGBColor(37, 99, 235)  # #2563EB
    C_BLUE_LIGHT = RGBColor(239, 246, 255)  # #EFF6FF
    C_TEXT_DARK = RGBColor(15, 23, 42)      # #0F172A
    C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
    C_GREEN = RGBColor(16, 185, 129)        # #10B981
    C_AMBER = RGBColor(245, 158, 11)        # #F59E0B
    C_WHITE = RGBColor(255, 255, 255)
    C_BG_PAGE = RGBColor(248, 250, 252)     # #F8FAFC
    C_CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0

    def add_bg(slide, color=C_BG_PAGE):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text):
        # Tag
        tb_tag = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        p_tag = tb_tag.text_frame.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = C_BLUE_PRIMARY

        # Title
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.55))
        p_title = tb_title.text_frame.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_DARK

        # Subtitle
        tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.4))
        p_sub = tb_sub.text_frame.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 1: COVER
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1, C_NAVY_DARK)

    # Accent top bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_BLUE_PRIMARY
    bar.line.fill.background()

    # Title box
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "BÁO CÁO TIẾN ĐỘ & ĐÁNH GIÁ CHẤM THI ĐỒ ÁN"
    p0.font.name = "Arial"
    p0.font.size = Pt(15)
    p0.font.bold = True
    p0.font.color.rgb = C_BLUE_PRIMARY

    p1 = tf.add_paragraph()
    p1.text = "HỆ THỐNG ĐĂNG KÝ HỌC PHẦN TRỰC TUYẾN"
    p1.font.name = "Arial"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE

    p2 = tf.add_paragraph()
    p2.text = "Đề tài số 15 • Kiến trúc Clean Architecture phân lớp chuẩn (mô hình eShop) • ASP.NET Core Blazor (.NET 8) & SQL Server"
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(148, 163, 184)

    # 4 Highlights Cards
    card_data = [
        ("4 PROJECT", "Clean Architecture", "CoreBusiness độc lập 0 ref", C_BLUE_PRIMARY),
        ("9 BẢNG 3NF", "SQL Server", "Header-Line 1 Transaction", C_GREEN),
        ("2 LUẬT CHÍNH", "CoreBusiness (6.0đ)", "Tiên quyết & Trùng lịch", C_AMBER),
        ("16 TIÊU CHÍ", "K1.1 ĐẾN K5.1", "Đạt 100% điểm tối đa", RGBColor(168, 85, 247))
    ]

    card_w = Inches(2.65)
    card_h = Inches(2.2)
    start_x = Inches(1.0)
    gap = Inches(0.24)

    for i, (head, sub, desc, acc_col) in enumerate(card_data):
        cx = start_x + i * (card_w + gap)
        cy = Inches(3.8)

        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = C_NAVY_CARD
        card.line.color.rgb = acc_col
        card.line.width = Pt(1.5)

        ctb = s1.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.2), card_w - Inches(0.4), card_h - Inches(0.4))
        ctf = ctb.text_frame
        ctf.word_wrap = True

        cp0 = ctf.paragraphs[0]
        cp0.text = head
        cp0.font.name = "Arial"
        cp0.font.size = Pt(18)
        cp0.font.bold = True
        cp0.font.color.rgb = acc_col

        cp1 = ctf.add_paragraph()
        cp1.text = sub
        cp1.font.name = "Arial"
        cp1.font.size = Pt(13)
        cp1.font.bold = True
        cp1.font.color.rgb = C_WHITE

        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.name = "Arial"
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = RGBColor(148, 163, 184)

    # Footer author info
    foot_tb = s1.shapes.add_textbox(Inches(1.0), Inches(6.4), Inches(11.3), Inches(0.6))
    fp = foot_tb.text_frame.paragraphs[0]
    fp.text = "Sinh viên thực hiện: Nguyễn Thị Lưu  •  Ngành: Tin học Kinh tế  •  Khoa: Hệ thống Thông tin Kinh tế  •  Trường ĐH Kinh tế – Đại học Huế"
    fp.font.name = "Arial"
    fp.font.size = Pt(11)
    fp.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # SLIDE 2: BẢNG TỔNG HỢP 16 TIÊU CHÍ (K1.1 -> K5.1)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "TỔNG QUAN ĐÁNH GIÁ CHẤM THI", "BẢNG TRA CỨU 16 TIÊU CHÍ CHẤM THI (K1.1 ĐẾN K5.1)", "Hệ thống đáp ứng trọn vẹn 100% tất cả các tiêu chí kỹ thuật và nghiệp vụ của đề tài")

    rows = 17
    cols = 5
    table_x = Inches(0.8)
    table_y = Inches(1.9)
    table_w = Inches(11.733)
    table_h = Inches(5.1)

    table_shape = s2.shapes.add_table(rows, cols, table_x, table_y, table_w, table_h)
    tbl = table_shape.table

    tbl.columns[0].width = Inches(1.1)
    tbl.columns[1].width = Inches(4.3)
    tbl.columns[2].width = Inches(2.9)
    tbl.columns[3].width = Inches(2.1)
    tbl.columns[4].width = Inches(1.333)

    headers = ["Mã", "Tiêu chí kiểm tra", "Cách kiểm tra của Giảng viên", "Vị trí trong mã nguồn", "Tối đa"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_BLUE_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER if j in [0, 4] else PP_ALIGN.LEFT

    criteria_summary = [
        ("K1.1", "Đủ 4 project; CoreBusiness (0 ref); UseCases chỉ ref CoreBusiness", "Mở file .csproj, xem ProjectReference", "4 project .csproj", "8.0 đ"),
        ("K1.2", "Interface repository ở UseCases, Plugin cài đặt; Razor inject interface", "Tìm @inject trong .razor; Program.cs", "UseCases / Program.cs", "7.0 đ"),
        ("K1.3", "2 luật nghiệp vụ cài trong CoreBusiness / UseCases, không nằm trong .razor", "Mở code EnrollmentValidationService.cs", "CoreBusiness.Services", "6.0 đ"),
        ("K1.4", "DI đăng ký đủ; giỏ / lựa chọn tạm dùng Scoped", "Xem Program.cs (Dòng 57-60)", "Program.cs", "4.0 đ"),
        ("K2.1", "Từ 5 bảng, 3NF, có khóa ngoại; script tạo bảng + dữ liệu chạy được", "Chạy DangKyTinChi_Database.sql trên SSMS", "Thư mục gốc .sql", "8.0 đ"),
        ("K2.2", "Cặp header–line ghi trong 1 transaction", "Code BeginTransactionAsync trong Plugin", "EnrollmentRepository.cs", "6.0 đ"),
        ("K2.3", "Mọi truy vấn tham số hóa, không nối chuỗi SQL", "Tìm $\"SELECT, + \"WHERE trong Plugin", "Plugin DataStore", "5.0 đ"),
        ("K2.4", "Luồng giao dịch chính chạy end-to-end, dữ liệu lưu SQL Server", "Thêm giỏ -> Gửi phiếu -> Sĩ số tự tăng 1", "Web -> SQL Server", "6.0 đ"),
        ("K3.1", "Ít nhất 3 component tái sử dụng có [Parameter] / EventCallback", "Mở Components/Shared", "SectionCard, Table...", "6.0 đ"),
        ("K3.2", "EditForm + DataAnnotations; lỗi nhập liệu hiển thị cạnh ô nhập", "Demo nhập sai tại Login & Quản trị", "Login / ManageCourses", "5.0 đ"),
        ("K3.3", "Ca vi phạm 2 luật nghiệp vụ hiện thông báo thân thiện, không crash", "Demo chọn môn thiếu tiên quyết & trùng lịch", "ValidationAlertComponent", "4.0 đ"),
        ("K3.4", "Dùng được trên màn hình điện thoại DevTools, bề rộng 390px", "Mở F12 -> iPhone 390px xem menu & TKB", "NavMenu / ScheduleTable", "5.0 đ"),
        ("K4.1", "Đăng nhập, đăng xuất; ít nhất 2 vai trò", "Demo 2 role: Student (sv21001) & Admin", "LoginComponent", "4.0 đ"),
        ("K4.2", "Tài khoản thường gõ thẳng URL trang quản trị bị chặn", "Gõ /admin/courses -> Bị chặn 403 Forbidden", "AccessDeniedComponent", "5.0 đ"),
        ("K4.3", "Nâng cao: Trợ lý AI và Xuất báo cáo (Excel/PDF) từ dữ liệu thật", "Demo Cố vấn AI và Xuất file PDF/Excel", "ReportExportService", "6.0 đ"),
        ("K5.1", "Use Case, Class, Sequence (PlantUML) khớp tên lớp và luồng code", "So sánh sơ đồ với kiến trúc code thực tế", "Sơ đồ Draw.io & PUML", "Tối đa")
    ]

    for idx, (code, name, test, loc, score) in enumerate(criteria_summary):
        row_idx = idx + 1
        data = [code, name, test, loc, score]
        for col_idx, val in enumerate(data):
            cell = tbl.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_WHITE if row_idx % 2 == 1 else RGBColor(241, 245, 249)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = "Arial"
            p.font.size = Pt(8.5)
            if col_idx in [0, 4]:
                p.alignment = PP_ALIGN.CENTER
                p.font.bold = True
                p.font.color.rgb = C_BLUE_PRIMARY if col_idx == 0 else C_GREEN
            else:
                p.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 3: NHÓM K1 (K1.1, K1.2, K1.4)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "NHÓM TIÊU CHÍ K1 • KIẾN TRÚC PHẦN MỀM (19.0 ĐIỂM)", "KIẾN TRÚC CLEAN ARCHITECTURE & DEPENDENCY INJECTION", "Tổ chức mã nguồn theo 4 project phân lớp chuẩn eShop, đảo ngược phụ thuộc (DIP) nghiêm ngặt")

    # Card Left
    cw = Inches(6.0)
    ch = Inches(4.9)
    cx = Inches(0.8)
    cy = Inches(1.9)

    card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch)
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_CARD_BORDER

    tb = s3.shapes.add_textbox(cx + Inches(0.3), cy + Inches(0.2), cw - Inches(0.6), ch - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "K1.1 • Đủ 4 Project & Quy tắc phụ thuộc (8.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets1 = [
        "CoreBusiness.csproj: Hoàn toàn 0 ProjectReference (Lõi độc lập tuyệt đối).",
        "UseCases.csproj: Dòng 10-12 chỉ tham chiếu duy nhất CoreBusiness.",
        "DataStore.SqlServer & Web: Phụ thuộc hướng tâm vào UseCases và CoreBusiness, không bị vòng lặp."
    ]
    for b in bullets1:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    p = tf.add_paragraph()
    p.text = "\nK1.2 • Interface ở UseCases, Plugin cài đặt (7.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets2 = [
        "Interface đặt tại UseCases/PluginInterfaces/DataStore/ (ICourseRepository, ISectionRepository...)",
        "Plugin DataStore.SqlServer cài đặt Interface bằng EF Core DbContext.",
        "File .razor 100% chỉ @inject Interface, tuyệt đối không phụ thuộc lớp cụ thể."
    ]
    for b in bullets2:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    p = tf.add_paragraph()
    p.text = "\nK1.4 • Cấu hình DI Container & Giỏ Scoped (4.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    p = tf.add_paragraph()
    p.text = "• Program.cs (Dòng 57-60): Đăng ký AddScoped cho IRegistrationCartStateStore và ICurrentUserStateStore, đảm bảo giỏ hàng lưu phiên làm việc độc lập của từng sinh viên."
    p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    # Image Right
    img_path1 = "report_images/image1.png"
    if os.path.exists(img_path1):
        s3.shapes.add_picture(img_path1, Inches(7.1), Inches(1.9), width=Inches(5.4))

    # ==========================================
    # SLIDE 4: TIÊU CHÍ K1.3 & K3.3 (2 LUẬT NGHIỆP VỤ)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "TRỌNG TÂM ĐỒ ÁN • TIÊU CHÍ K1.3 (6.0Đ) & K3.3 (4.0Đ)", "CÀI ĐẶT 2 LUẬT NGHIỆP VỤ & XỬ LÝ VI PHẠM THÂN THIỆN", "Luật Tiên Quyết và Chống Trùng Lịch được đóng gói trong CoreBusiness, hiển thị cảnh báo đẹp mắt")

    # Card Left
    card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch)
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_CARD_BORDER

    tb = s4.shapes.add_textbox(cx + Inches(0.3), cy + Inches(0.2), cw - Inches(0.6), ch - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "K1.3 • Luật 1: Kiểm tra Học phần Tiên quyết (3.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "Vị trí code: EnrollmentValidationService.cs (Dòng 26 - 43).",
        "Thuật toán: Đối chiếu danh sách môn tiên quyết trong bảng Prerequisites với lịch sử tích lũy (CompletedCourses) của sinh viên (điểm >= 4.0).",
        "Kịch bản Demo: Sinh viên sv21002 (Trần Thị Mai) chưa học môn OOP (TIN102). Khi chọn môn Web (WEB301), hệ thống phát hiện và chặn ghi phiếu."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    p = tf.add_paragraph()
    p.text = "\nK1.3 • Luật 2: Kiểm tra Chống Trùng Lịch học (3.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "Vị trí code: EnrollmentValidationService.cs (Dòng 46 - 70).",
        "Thuật toán: Lặp qua các buổi học (SectionSchedules) trong tuần, phát hiện trùng khi cùng Thứ (DayOfWeek) và giao nhau khoảng tiết (StartPeriod <= other.EndPeriod && EndPeriod >= other.StartPeriod).",
        "Kịch bản Demo: Sinh viên sv21001 chọn lớp WEB301_01 (Thứ 2 tiết 1-4) và TIN201_01 (Thứ 2 tiết 1-3). Hệ thống từ chối giao dịch."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    p = tf.add_paragraph()
    p.text = "\nK3.3 • Thông báo vi phạm thân thiện (4.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_GREEN

    p = tf.add_paragraph()
    p.text = "• Không bao giờ văng Exception/Crash. Hiển thị alert qua ValidationAlertComponent giải thích rõ: tên môn tiên quyết còn thiếu / mã lớp và khoảng tiết bị trùng."
    p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    # Image Right
    img_path2 = "report_images/image2.png"
    if os.path.exists(img_path2):
        s4.shapes.add_picture(img_path2, Inches(7.1), Inches(1.9), width=Inches(5.4))

    # ==========================================
    # SLIDE 5: NHÓM K2 (CSDL & TRANSACTION)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "NHÓM TIÊU CHÍ K2 • CƠ SỞ DỮ LIỆU & TRANSACTION (25.0 ĐIỂM)", "9 BẢNG 3NF & GIAO DỊCH HEADER - LINE AN TOÀN TUYỆT ĐỐI", "Dữ liệu lưu trữ trong Microsoft SQL Server, đảm bảo tính toàn vẹn ACID qua 1 Transaction")

    card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch)
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_CARD_BORDER

    tb = s5.shapes.add_textbox(cx + Inches(0.3), cy + Inches(0.2), cw - Inches(0.6), ch - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "K2.1 • CSDL 9 Bảng Chuẩn 3NF (8.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "File DangKyTinChi_Database.sql chạy thành công 100% trên CSDL trống.",
        "9 bảng có khóa chính, khóa ngoại đầy đủ: Courses, Students, Prerequisites, Sections, SectionSchedules, CompletedCourses, Enrollments, EnrollmentLines, UserAccounts."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    p = tf.add_paragraph()
    p.text = "\nK2.2 • Cặp Header–Line Ghi Trong 1 Transaction (6.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "Vị trí code: EnrollmentRepository.cs (Dòng 83 - 114).",
        "Sử dụng context.Database.BeginTransactionAsync() tường minh.",
        "Ghi đồng thời bảng Enrollments (Header) và EnrollmentLines (Line). Nếu có lỗi tự động Rollback, không bao giờ để lại dòng đăng ký mồ côi."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    p = tf.add_paragraph()
    p.text = "\nK2.3 & K2.4 • Tham số hóa 100% & Giao dịch End-to-End (11.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "100% truy vấn dùng LINQ tham số hóa, không nối chuỗi SQL thô (chống SQL Injection).",
        "Luồng End-to-End: Chọn môn -> Đặt phiếu -> Sĩ số CurrentEnrollment của lớp tự động tăng 1; Hủy môn sau hạn: Tự động trừ 1 nhả chỗ."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    # Image Right
    img_path3 = "report_images/image3.png"
    if os.path.exists(img_path3):
        s5.shapes.add_picture(img_path3, Inches(7.1), Inches(1.9), width=Inches(5.4))

    # ==========================================
    # SLIDE 6: NHÓM K3 (COMPONENT & MOBILE 390PX)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "NHÓM TIÊU CHÍ K3 • GIAO DIỆN & TỐI ƯU TRẢI NGHIỆM (16.0 ĐIỂM)", "COMPONENT TÁI SỬ DỤNG, FORM VALIDATION & MOBILE 390PX", "Thiết kế Responsive DevTools 390px, giao tiếp EventCallback 2 chiều, báo lỗi nhập liệu tức thì")

    card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch)
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_CARD_BORDER

    tb = s6.shapes.add_textbox(cx + Inches(0.3), cy + Inches(0.2), cw - Inches(0.6), ch - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "K3.1 • 3 Component Tái Sử Dụng Chuẩn Blazor (6.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "1. SectionCardComponent: [Parameter] Section, IsInCart; EventCallback OnAddToCart, OnRemoveFromCart.",
        "2. ValidationAlertComponent: [Parameter] Errors; EventCallback OnDismiss.",
        "3. ScheduleTableComponent: [Parameter] Sections, EventCallback<Section> OnSectionClick."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    p = tf.add_paragraph()
    p.text = "\nK3.2 • EditForm + DataAnnotations (5.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "Áp dụng tại LoginComponent, ManageCoursesComponent, ManageSectionsComponent.",
        "Dùng <DataAnnotationsValidator /> và <ValidationMessage For=...> hiển thị lỗi đỏ ngay cạnh ô nhập liệu khi người dùng nhập sai."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    p = tf.add_paragraph()
    p.text = "\nK3.4 • Dùng Được Trên Màn Hình Điện Thoại 390px (5.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "Mở F12 DevTools -> Chọn iPhone 12/14 Pro (Width: 390px).",
        "Thanh menu co thành nút hamburger 3 gạch thu gọn.",
        "Thời khóa biểu có chế độ Agenda xem theo ngày cuộn ngang, tối ưu hoàn hảo cho màn hình cảm ứng di động."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    # Image Right
    img_path6 = "report_images/image6.png"
    if os.path.exists(img_path6):
        s6.shapes.add_picture(img_path6, Inches(7.1), Inches(1.9), width=Inches(5.4))

    # ==========================================
    # SLIDE 7: NHÓM K4 (BẢO MẬT & NÂNG CAO)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "NHÓM TIÊU CHÍ K4 • BẢO MẬT & TÍNH NĂNG NÂNG CAO (15.0 ĐIỂM)", "PHÂN QUYỀN 2 VAI TRÒ, CHẶN URL 403, TRỢ LÝ AI & XUẤT BÁO CÁO", "Bảo mật kiểm soát truy cập phân hệ Quản trị, tích hợp AI Cố vấn và xuất báo cáo PDF/Excel từ CSDL thật")

    card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, cy, cw, ch)
    card.fill.solid()
    card.fill.fore_color.rgb = C_WHITE
    card.line.color.rgb = C_CARD_BORDER

    tb = s7.shapes.add_textbox(cx + Inches(0.3), cy + Inches(0.2), cw - Inches(0.6), ch - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "K4.1 & K4.2 • Phân Quyền & Chặn URL Trang Quản Trị (9.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "2 Vai trò rõ rệt: Sinh viên (sv21001-sv21007) và Quản trị viên Admin (admin/admin123).",
        "Có nút Đăng xuất xóa phiên làm việc tức thì.",
        "Chặn URL trực tiếp: Sinh viên hoặc khách chưa đăng nhập gõ thẳng /admin/courses hoặc /admin/enrollments lập tức bị chặn và chuyển về trang /access-denied (403 Forbidden)."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    p = tf.add_paragraph()
    p.text = "\nK4.3 • Tính Năng Nâng Cao: AI & Xuất Báo Cáo (6.0đ)"
    p.font.name = "Arial"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    bullets = [
        "1. Trợ lý Cố vấn AI (AiAdvisor & AiAdminAssistant): Chatbot thông minh tư vấn học tập, kiểm tra điều kiện tiên quyết và gợi ý lộ trình môn học cho sinh viên.",
        "2. Xuất báo cáo PDF & Excel (ReportExportService): Sử dụng QuestPDF và ClosedXML xuất trực tiếp Danh sách lớp, Thời khóa biểu và Phiếu đăng ký từ CSDL SQL Server thật ra file .pdf và .xlsx theo mẫu chuẩn."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.name = "Arial"; p.font.size = Pt(10.5); p.font.color.rgb = C_TEXT_DARK

    # Image Right
    img_path9 = "report_images/image9.png"
    if os.path.exists(img_path9):
        s7.shapes.add_picture(img_path9, Inches(7.1), Inches(1.9), width=Inches(5.4))

    # ==========================================
    # SLIDE 8: NHÓM K5 (SƠ ĐỒ THIẾT KẾ HỆ THỐNG)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "NHÓM TIÊU CHÍ K5 • THIẾT KẾ HỆ THỐNG (ĐIỂM TỐI ĐA)", "SƠ ĐỒ USE CASE & CLASS DIAGRAM KHỚP 100% VỚI MÃ NGUỒN", "Thiết kế chuẩn hóa bằng PlantUML và Draw.io, đồng nhất tuyệt đối giữa phân tích và hiện thực")

    # 2 image columns
    col_w = Inches(5.6)
    col_h = Inches(4.8)

    # Use Case Diagram
    uc_path = "DangKiTinChi-Usecase.drawio.png"
    if os.path.exists(uc_path):
        s8.shapes.add_picture(uc_path, Inches(0.8), Inches(2.0), width=col_w)

    # Class Diagram
    cl_path = "DangKiTinChi-Sơ đồ Lớp .drawio.png"
    if os.path.exists(cl_path):
        s8.shapes.add_picture(cl_path, Inches(6.8), Inches(2.0), width=col_w)

    # Sub-captions
    tb_c1 = s8.shapes.add_textbox(Inches(0.8), Inches(1.7), col_w, Inches(0.3))
    p = tb_c1.text_frame.paragraphs[0]
    p.text = "▲ Sơ đồ Use Case: 2 Actor Sinh viên & Cán bộ, 3 Phân hệ, Include/Extend chuẩn"
    p.font.name = "Arial"; p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    tb_c2 = s8.shapes.add_textbox(Inches(6.8), Inches(1.7), col_w, Inches(0.3))
    p = tb_c2.text_frame.paragraphs[0]
    p.text = "▲ Sơ đồ Lớp (Class Diagram): 9 Entity 3NF, Cặp Header-Line, Service kiểm tra luật"
    p.font.name = "Arial"; p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

    # ==========================================
    # SLIDE 9: KỊCH BẢN DEMO 3 PHÚT CHO GIẢNG VIÊN
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_bg(s9)
    add_header(s9, "HƯỚNG DẪN THỰC HÀNH", "KỊCH BẢN DEMO 3 PHÚT DÀNH CHO GIẢNG VIÊN CHẤM THI", "Quy trình demo trực quan, dễ dàng thẩm định 16 tiêu chí chấm thi trong thời gian ngắn nhất")

    demo_steps = [
        ("BƯỚC 1: DEMO 2 LUẬT NGHIỆP VỤ (K1.3, K3.3)", 
         "• Đăng nhập sv21002 (chưa học OOP) -> Đăng ký môn Web (WEB301) -> Bị chặn Tiên quyết.\n• Đăng nhập sv21001 -> Chọn lớp WEB301_01 (Thứ 2 tiết 1-4) & TIN201_01 (Thứ 2 tiết 1-3) -> Bị chặn Trùng lịch."),
        ("BƯỚC 2: DEMO PHÂN QUYỀN & CHẶN URL (K4.1, K4.2)", 
         "• Khi đang ở quyền Sinh viên, gõ trực tiếp URL /admin/courses trên thanh địa chỉ.\n• Hệ thống chặn ngay lập tức và điều hướng về trang HTTP 403 Access Denied."),
        ("BƯỚC 3: DEMO TRANSACTION HEADER-LINE & HỦY MÔN (K2.2, K2.4)", 
         "• Sinh viên gửi phiếu đăng ký -> Dữ liệu ghi vào Enrollments & EnrollmentLines trong 1 Transaction, sĩ số tăng 1.\n• Đăng nhập admin -> Vào Rút môn cho SV -> Sĩ số tự động trừ 1 nhả chỗ."),
        ("BƯỚC 4: DEMO XUẤT FILE EXCEL / PDF & CỐ VẤN AI (K4.3)", 
         "• Bấm nút 'Xuất Excel' hoặc 'Xuất PDF' tại trang danh sách lớp -> File chuẩn tải về máy.\n• Vào mục 'Cố vấn AI' trò chuyện nhận tư vấn học phần theo học kỳ.")
    ]

    dw = Inches(5.6)
    dh = Inches(2.3)
    coords = [
        (Inches(0.8), Inches(1.9)),
        (Inches(6.8), Inches(1.9)),
        (Inches(0.8), Inches(4.5)),
        (Inches(6.8), Inches(4.5))
    ]

    for (step_title, step_desc), (dx, dy) in zip(demo_steps, coords):
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, dx, dy, dw, dh)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BLUE_PRIMARY
        card.line.width = Pt(1.5)

        tb = s9.shapes.add_textbox(dx + Inches(0.2), dy + Inches(0.15), dw - Inches(0.4), dh - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = step_title
        p.font.name = "Arial"; p.font.size = Pt(11.5); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY

        p = tf.add_paragraph()
        p.text = step_desc
        p.font.name = "Arial"; p.font.size = Pt(10); p.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 10: CONCLUSION & THANKS
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_bg(s10, C_NAVY_DARK)

    bar = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_GREEN
    bar.line.fill.background()

    tb = s10.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.3), Inches(3.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "TỔNG KẾT BÁO CÁO TIẾN ĐỘ ĐỒ ÁN"
    p.font.name = "Arial"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = C_GREEN
    p.alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = "HOÀN THÀNH 100% TIÊU CHÍ (K1.1 ĐẾN K5.1)"
    p.font.name = "Arial"; p.font.size = Pt(32); p.font.bold = True; p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = "Đồ án Đăng ký Học phần Trực tuyến (Đề 15) tuân thủ nghiêm ngặt mô hình Clean Architecture,\nđầy đủ luật nghiệp vụ, giao dịch an toàn và tối ưu giao diện trải nghiệm người dùng."
    p.font.name = "Arial"; p.font.size = Pt(14); p.font.color.rgb = RGBColor(148, 163, 184)
    p.alignment = PP_ALIGN.CENTER

    p = tf.add_paragraph()
    p.text = "\nKính chúc Thầy/Cô và Hội đồng chấm thi sức khỏe và công tác tốt!\nXin chân thành cảm ơn!"
    p.font.name = "Arial"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = C_BLUE_PRIMARY
    p.alignment = PP_ALIGN.CENTER

    prs.save("BaoCao_TienDo_DangKyHocPhan_De15.pptx")
    print("Presentation saved successfully: BaoCao_TienDo_DangKyHocPhan_De15.pptx")

if __name__ == "__main__":
    create_deck()
